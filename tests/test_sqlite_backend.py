"""SQLite 后端集成冒烟测试（阶段 D 前置验证）。

在真实 `create_app()` 上用 `STORAGE_BACKEND=sqlite` 启动（data.db/offline.db 重定向到
工作区可写临时目录，不触碰真实数据库与 YAML），经 HTTP test_client 走真实路由，
验证 sqlite 仓储栈（含新 SqliteThreadRepository）在应用内可运行且数据确实落库。
测试结束后恢复配置——不入切换、不迁移、不改你的 YAML/真实 data.db。
"""

from __future__ import annotations

import sqlite3
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest

from lib.repositories.sqlite_thread import SqliteThreadRepository


@pytest.fixture(scope="module")
def sqlite_app(tmp_path_factory):
    import importlib.util
    import sys

    import lib.offline as offmod
    import lib.e2ee.flask_setup as e2ee

    tmp = tmp_path_factory.mktemp("sqlite-app")
    settings = tmp / "config.yaml"
    settings.write_text("{}", encoding="utf-8")
    source = Path(__file__).resolve().parent.parent / "app.py"
    with pytest.MonkeyPatch.context() as patch:
        patch.setenv("EXDIARY_SETTINGS", str(settings))
        patch.setenv("EXPERMATE_DATA_DIR", str(tmp))
        patch.setenv("EXDIARY_DB", str(tmp / "data.db"))
        patch.delenv("EXDIARY_ENC_KEY", raising=False)
        patch.setattr(offmod, "offline_db_path", lambda: tmp / "offline.db")
        patch.setattr(e2ee, "_ACCOUNT_DB", tmp / "_e2ee_accounts.db")
        patch.setattr(e2ee, "_KMS_KEY_FILE", tmp / "_e2ee_kms.key")
        spec = importlib.util.spec_from_file_location("sqlite_smoke_app", source)
        appmod = importlib.util.module_from_spec(spec)
        patch.setitem(sys.modules, "sqlite_smoke_app", appmod)
        spec.loader.exec_module(appmod)
        app = appmod.create_app()
        app.config.update(TESTING=True, SMOKE_DATA_DIR=tmp)
        try:
            yield app
        finally:
            appmod.ctx.shutdown()
            security = app.extensions.get("e2ee")
            if security:
                security.conn.close()
                security.jconn.close()


def test_app_boots_in_sqlite_mode(sqlite_app):
    assert sqlite_app is not None
    # sqlite 分支应创建 offline.db（重定向到工作区临时目录）
    db = sqlite_app.config["SMOKE_DATA_DIR"] / "offline.db"
    assert db.exists()


def test_react_default_entry_and_explicit_offline_scope(sqlite_app, tmp_path):
    dist = tmp_path / "dist"
    dist.mkdir()
    (dist / "index.html").write_text("<html>React integration entry</html>", encoding="utf-8")
    sqlite_app.config["REACT_DIST_DIR"] = dist
    client = sqlite_app.test_client()
    for path in ("/", "/experiments", "/settings", "/analyze", "/analysis/report", "/app/"):
        assert b"React integration entry" in client.get(path).data
    # 通过真实应用装配选择独立离线库，而不是接口自行回退。
    assert client.get("/api/v1/experiments").status_code == 200
    assert client.get("/api/v1/settings").status_code == 200
    client.set_cookie("exdiary_token", "invalid-session", domain="localhost")
    assert client.get("/api/v1/experiments").status_code == 401
    assert client.get("/api/v1/experiments").status_code == 200


def test_offline_thread_repo_is_sqlite(sqlite_app):
    """无认证请求 → g.thread_repo 应为 SqliteThreadRepository（线程已进 SQLite）。"""
    with sqlite_app.test_request_context("/"):
        sqlite_app.preprocess_request()
        from flask import g
        assert isinstance(g.thread_repo, SqliteThreadRepository)
        # 其它仓储也应是 sqlite 实现
        from lib.repositories.sqlite_experiment import SqliteExperimentRepository
        assert isinstance(g.exp_repo, SqliteExperimentRepository)


def test_cookie_login_reads_account_chat_state(sqlite_app):
    """桌面端仅靠持久 Cookie 也必须读取登录账号的聊天状态，而非离线库。"""
    from flask import g
    from lib.auth import create_token

    token = create_token("desktop-user", 1)
    with sqlite_app.test_request_context("/", headers={"Cookie": f"exdiary_token={token}"}):
        sqlite_app.preprocess_request()
        assert g.user_id == "desktop-user"
        g.thread_repo.save_current_state({
            "history": [{"role": "user", "content": "账号聊天记录"}],
        })

    client = sqlite_app.test_client()
    client.set_cookie("exdiary_token", token, domain="localhost")
    response = client.get("/api/agent/history")
    assert response.status_code == 200
    assert response.get_json()["history"][0]["content"] == "账号聊天记录"


def test_routes_serve_and_persist(sqlite_app):
    c = sqlite_app.test_client()

    # 1. 列实验（空库 → 200）
    rv = c.get("/api/v1/experiments")
    assert rv.status_code == 200

    # 2. 保存一条实验（走 exp_repo + update_log_repo，无 LLM）
    rv = c.post("/experiments/EXP-2026-001/save-json", json={
        "id": "EXP-2026-001", "title": "SQLite冒烟实验", "date": "2026-08-19",
        "experimenter": "测试", "status": "planned", "tags": ["smoke"],
        "purpose": "验证 sqlite 后端", "original_notes": "",
        "materials": [], "results": {},
    })
    assert rv.status_code == 200, rv.get_json()
    assert rv.get_json()["ok"] is True

    # 3. 读回该实验（走 sqlite exp_repo）
    rv = c.get("/experiments/EXP-2026-001/yaml")
    assert rv.status_code == 200

    # 4. 收藏：建收藏夹 + 置顶 + 列表（走 sqlite favorites；前缀为 /api）
    rv = c.post("/api/collections", json={"name": "冒烟收藏夹"})
    assert rv.status_code == 200 and rv.get_json()["ok"] is True
    rv = c.post("/api/experiments/EXP-2026-001/pin", json={})
    assert rv.status_code == 200 and rv.get_json()["pinned"] is True
    rv = c.get("/api/list-collections")
    assert rv.status_code == 200

    # 5. 分析历史（sqlite analyses，空）
    rv = c.get("/api/analysis-history")
    assert rv.status_code == 200

    # 6. 搜索（FTS）
    rv = c.get("/api/experiments/search", query_string={"q": "冒烟"})
    assert rv.status_code == 200


def test_chat_history_search_routes(sqlite_app):
    """搜索与附近上下文接口只返回小窗口，且能跨归档读取。"""
    with sqlite_app.test_request_context("/"):
        sqlite_app.preprocess_request()
        from flask import g
        g.thread_repo.archive_compressed_history("SES-smoke", 0, [
            {"role": "user", "content": "需要找回旧的电池讨论",
             "created_at": "2026-08-20T09:00:00Z"},
        ])

    client = sqlite_app.test_client()
    rv = client.get("/api/agent/history/search", query_string={"q": "电池"})
    assert rv.status_code == 200
    match = rv.get_json()["matches"][0]
    assert match["session_id"] == "SES-smoke"

    rv = client.get("/api/agent/history/context", query_string={
        "session_id": match["session_id"], "sequence": match["sequence"],
    })
    assert rv.status_code == 200
    assert rv.get_json()["messages"][0]["content"] == "需要找回旧的电池讨论"
    rv = client.get("/api/agent/history/sessions")
    assert rv.status_code == 200
    assert rv.get_json()["sessions"][0]["session_id"] == "SES-smoke"
    rv = client.get("/api/agent/history/browse", query_string={"session_id": "SES-smoke"})
    assert rv.status_code == 200
    assert rv.get_json()["messages"][0]["content"] == "需要找回旧的电池讨论"


def test_manual_edit_rejects_stale_revision(sqlite_app):
    client = sqlite_app.test_client()
    exp_id = "EXP-2026-REVISION"
    first = {
        "id": exp_id, "title": "版本测试", "date": "2026-08-21",
        "experimenter": "测试", "status": "planned", "tags": [], "purpose": "初始值",
    }
    assert client.post(f"/experiments/{exp_id}/save-json", json=first).get_json()["revision"] == 1
    second = dict(first, purpose="窗口 A 保存", expected_revision=1)
    assert client.post(f"/experiments/{exp_id}/save-json", json=second).get_json()["revision"] == 2
    stale = dict(first, purpose="窗口 B 旧页面", expected_revision=1)
    rv = client.post(f"/experiments/{exp_id}/save-json", json=stale)
    assert rv.status_code == 409
    assert rv.get_json()["revision"] == 2


def test_attachment_upload_link_download_and_remove(sqlite_app):
    client = sqlite_app.test_client()
    exp_id = "EXP-2026-ATT"
    rv = client.post(f"/experiments/{exp_id}/save-json", json={
        "id": exp_id, "title": "附件测试", "date": "2026-08-21",
        "experimenter": "测试", "status": "planned", "tags": [],
    })
    assert rv.status_code == 200, rv.get_json()

    rv = client.post("/api/attachments", data={
        "experiment_id": exp_id, "title": "原始数据",
        "file": (BytesIO(b"time,value\\n0,1\\n"), "result.csv"),
    }, content_type="multipart/form-data")
    assert rv.status_code == 200, rv.get_json()
    attachment = rv.get_json()["attachment"]
    sha = attachment["sha256"]

    rv = client.get(f"/api/experiments/{exp_id}/attachments")
    assert rv.status_code == 200
    assert rv.get_json()["attachments"][0]["title"] == "原始数据"
    assert client.get(f"/api/v1/experiments/{exp_id}").status_code == 200
    rv = client.get(f"/api/attachments/{sha}?download=1")
    assert rv.status_code == 200
    assert rv.data == b"time,value\\n0,1\\n"

    rv = client.delete(f"/api/experiments/{exp_id}/attachments/{sha}")
    assert rv.status_code == 200
    assert client.get(f"/api/experiments/{exp_id}/attachments").get_json()["attachments"] == []


def test_chat_attachment_can_be_managed_by_agent(sqlite_app):
    """聊天上传不自动关联实验；Agent 工具可搜索、读取并在语义确认后关联。"""
    client = sqlite_app.test_client()
    exp_id = "EXP-2026-CHAT-ATT"
    client.post(f"/experiments/{exp_id}/save-json", json={
        "id": exp_id, "title": "聊天附件测试", "date": "2026-08-21",
        "experimenter": "测试", "status": "planned", "tags": [],
    })
    rv = client.post("/api/agent/attachments", data={
        "file": (BytesIO(b"time,value\n0,1\n"), "chat-data.csv"),
    }, content_type="multipart/form-data")
    assert rv.status_code == 200, rv.get_json()
    attachment = rv.get_json()["attachment"]
    assert client.get(f"/api/experiments/{exp_id}/attachments").get_json()["attachments"] == []

    with sqlite_app.test_request_context("/"):
        sqlite_app.preprocess_request()
        from flask import g
        from lib.agent_v2 import ToolExecutor
        tools = ToolExecutor(g.exp_repo, attachment_store=g.attachment_store)
        exact = tools.execute("read_experiment", {
            "refs": [exp_id], "fields": ["title", "attachments", "revision"],
        }, None)
        assert exact["experiments"][exp_id]["title"] == "聊天附件测试"
        found = tools.execute("search_attachments", {"query": "chat-data"}, None)
        assert found["attachments"][0]["sha256"] == attachment["sha256"]
        read = tools.execute("read_attachment", {"sha256": attachment["sha256"]}, None)
        assert "time,value" in read["content"]
        linked = tools.execute("manage_attachment", {
            "action": "link", "sha256": attachment["sha256"], "exp_id": exp_id,
            "title": "时间序列原始数据", "description": "聊天中上传，供本实验记录使用",
        }, None)
        assert linked["status"] == "ok"
        assert g.exp_repo.load(exp_id)["attachments"][0]["title"] == "时间序列原始数据"


def test_xlsx_attachment_has_native_preview(sqlite_app):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    book = Workbook()
    sheet = book.active
    sheet.title = "数据"
    sheet.append(["时间", "数值"])
    sheet.append(["10:00", 1.234567])
    sheet.merge_cells("A3:B3")
    sheet["A3"] = "合并单元格"
    sheet.column_dimensions["A"].width = 20
    sheet.row_dimensions[1].height = 24
    sheet["A1"].font = Font(name="Arial", sz=15, bold=True, color="FF0000")
    sheet["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet["A1"].border = Border(left=Side(style="medium", color="0000FF"))
    raw = BytesIO()
    book.save(raw)

    client = sqlite_app.test_client()
    rv = client.post("/api/agent/attachments", data={
        "file": (BytesIO(raw.getvalue()), "native-preview.xlsx"),
    }, content_type="multipart/form-data")
    assert rv.status_code == 200, rv.get_json()
    sha = rv.get_json()["attachment"]["sha256"]
    preview = client.get(f"/api/attachments/{sha}/preview")
    assert preview.status_code == 200, preview.get_json()
    data = preview.get_json()
    assert data["kind"] == "xlsx"
    assert data["sheet"] == "数据"
    assert data["rows"][0] == ["时间", "数值"]
    assert data["rows"][1] == ["10:00", "1.2346"]
    assert data["column_widths"]["A"] > data["column_widths"]["B"]
    assert data["row_heights"]["1"] == 32
    assert data["merged_cells"] == [{"start_row": 3, "start_col": 1, "end_row": 3, "end_col": 2}]
    assert data["cell_styles"]["A1"] == {
        "font": {"family": "Arial", "size": 20, "bold": True, "color": "#FF0000"},
        "fill": "#FFFF00",
        "alignment": {"horizontal": "center", "vertical": "middle", "wrap": True},
        "borders": {"left": "2px solid #0000FF"},
    }


def test_read_xlsx_attachment_truncates_each_sheet_independently(sqlite_app):
    from openpyxl import Workbook

    book = Workbook()
    first = book.active
    first.title = "前面的大表"
    first.append(["第一张表很长" * 30])
    second = book.create_sheet("后面的表")
    second.append(["后续工作表仍可见", 42])
    raw = BytesIO()
    book.save(raw)

    client = sqlite_app.test_client()
    uploaded = client.post("/api/agent/attachments", data={
        "file": (BytesIO(raw.getvalue()), "multi-sheet.xlsx"),
    }, content_type="multipart/form-data")
    assert uploaded.status_code == 200, uploaded.get_json()
    sha = uploaded.get_json()["attachment"]["sha256"]

    with sqlite_app.test_request_context("/"):
        sqlite_app.preprocess_request()
        from flask import g
        from lib.agent_v2 import ToolExecutor
        result = ToolExecutor(g.exp_repo, attachment_store=g.attachment_store).execute(
            "read_attachment", {"sha256": sha, "max_chars": 20}, None,
        )
        assert result["extraction"] == "xlsx"
        assert result["sheet_count"] == 2
        assert result["sheets"][0]["truncated"] is True
        assert result["sheets"][0]["truncation_reasons"] == ["chars"]
        assert result["sheets"][1]["content"] == "后续工作表仍可见\t42"

        continued = ToolExecutor(g.exp_repo, attachment_store=g.attachment_store).execute(
            "read_attachment", {"sha256": sha, "sheet": "后面的表", "start_row": 1}, None,
        )
        assert [sheet["name"] for sheet in continued["sheets"]] == ["后面的表"]


def test_read_image_attachment_uses_current_model_vision(sqlite_app):
    class VisionLLM:
        def __init__(self):
            self.calls = []

        def describe_image(self, content, mime, name):
            self.calls.append((content, mime, name))
            return "画面是一张黄色薄膜的实验照片。"

    client = sqlite_app.test_client()
    uploaded = client.post("/api/agent/attachments", data={
        "file": (BytesIO(b"not-used-by-the-test-model"), "film.jpg"),
    }, content_type="multipart/form-data")
    assert uploaded.status_code == 200, uploaded.get_json()
    attachment = uploaded.get_json()["attachment"]
    llm = VisionLLM()
    loop = SimpleNamespace(
        llm=llm,
        child=SimpleNamespace(agent_role=None, exp_id=None),
    )

    with sqlite_app.test_request_context("/"):
        sqlite_app.preprocess_request()
        from flask import g
        from lib.agent_v2 import ToolExecutor
        result = ToolExecutor(g.exp_repo, attachment_store=g.attachment_store).execute(
            "read_attachment", {"sha256": attachment["sha256"]}, loop,
        )

    assert result["status"] == "ok"
    assert result["extraction"] == "vision"
    assert result["content"] == "画面是一张黄色薄膜的实验照片。"
    assert llm.calls == [(b"not-used-by-the-test-model", "image/jpeg", "film.jpg")]


def test_data_persisted_to_sqlite(sqlite_app):
    """验证写入确实进了 offline.db（而非 YAML）。"""
    tmp = sqlite_app.config["SMOKE_DATA_DIR"]
    conn = sqlite3.connect(str(tmp / "offline.db"))
    n = conn.execute("SELECT COUNT(*) FROM experiments").fetchone()[0]
    t = [r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")]
    conn.close()
    assert n >= 1
    assert "threads" in t
    assert "kv_store" in t
